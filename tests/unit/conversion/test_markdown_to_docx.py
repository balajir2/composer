"""Tests for the Markdown -> DOCX renderer."""

import io
from typing import Any

from docx import Document  # pyright: ignore[reportMissingImports]


def _load(content: bytes) -> Any:
    # python-docx's top-level `Document` name is a factory *function*, not
    # the `Document` class it returns -- annotating the return type with it
    # directly trips pyright's reportGeneralTypeIssues. `Any` matches the
    # convention already used for python-docx call sites elsewhere in this
    # repo (see src/api/uploads.py, tests/unit/api/test_uploads.py), which
    # likewise leave docx objects untyped rather than fight its stubs.
    return Document(io.BytesIO(content))


def test_heading_becomes_docx_heading() -> None:
    from src.conversion.markdown_to_docx import markdown_to_docx

    doc = _load(markdown_to_docx("# Title\n\nSome text."))
    styles = [p.style.name for p in doc.paragraphs if p.text]
    assert any(s.startswith("Heading") for s in styles)
    texts = [p.text for p in doc.paragraphs]
    assert "Title" in texts
    assert "Some text." in texts


def test_bullet_list_becomes_list_paragraphs() -> None:
    from src.conversion.markdown_to_docx import markdown_to_docx

    doc = _load(markdown_to_docx("- one\n- two\n- three"))
    texts = [p.text for p in doc.paragraphs if p.text]
    assert texts == ["one", "two", "three"]


def test_table_becomes_docx_table() -> None:
    from src.conversion.markdown_to_docx import markdown_to_docx

    md = "| A | B |\n|---|---|\n| 1 | 2 |\n| 3 | 4 |"
    doc = _load(markdown_to_docx(md))
    assert len(doc.tables) == 1
    table = doc.tables[0]
    assert [c.text for c in table.rows[0].cells] == ["A", "B"]
    assert [c.text for c in table.rows[1].cells] == ["1", "2"]
    assert [c.text for c in table.rows[2].cells] == ["3", "4"]


def test_bold_and_italic_preserved_as_text() -> None:
    """Fidelity target: run-level bold/italic on simple inline spans, not a
    full round-trip of every possible markdown construct."""
    from src.conversion.markdown_to_docx import markdown_to_docx

    doc = _load(markdown_to_docx("**bold** and *italic*"))
    all_text = " ".join(p.text for p in doc.paragraphs)
    assert "bold" in all_text
    assert "italic" in all_text


def test_empty_content_produces_valid_empty_document() -> None:
    from src.conversion.markdown_to_docx import markdown_to_docx

    doc = _load(markdown_to_docx(""))
    assert doc is not None


def test_table_cell_markdown_is_parsed_not_literal() -> None:
    """Table cells must have markdown syntax stripped (or rendered as real
    bold/italic formatting) rather than leaking literal `**`/`*` characters
    into the cell text."""
    from src.conversion.markdown_to_docx import markdown_to_docx

    md = "| Priority | Note |\n|---|---|\n| **High** | *urgent* |"
    doc = _load(markdown_to_docx(md))
    table = doc.tables[0]
    cell_texts = [c.text for c in table.rows[1].cells]
    assert "High" in cell_texts[0]
    assert "*" not in cell_texts[0]
    assert "urgent" in cell_texts[1]
    assert "*" not in cell_texts[1]


def test_inline_code_span_preserved_in_paragraph() -> None:
    from src.conversion.markdown_to_docx import markdown_to_docx

    doc = _load(markdown_to_docx("Set the `status` field to `Active`."))
    all_text = " ".join(p.text for p in doc.paragraphs)
    assert "status" in all_text
    assert "Active" in all_text
    assert all_text == "Set the status field to Active."


def test_inline_code_span_preserved_in_table_cell() -> None:
    from src.conversion.markdown_to_docx import markdown_to_docx

    md = "| Field | Value |\n|---|---|\n| `status` | `Active` |"
    doc = _load(markdown_to_docx(md))
    table = doc.tables[0]
    cell_texts = [c.text for c in table.rows[1].cells]
    assert cell_texts == ["status", "Active"]


def test_bracket_placeholder_is_not_silently_dropped() -> None:
    """`<Client Name>`-style bracket placeholders (common in BRDs) must
    survive conversion as literal text, not vanish. With CommonMark's raw-
    HTML passthrough enabled (the pre-fix default), `<Client Name>` parses
    as an `html_inline` token that `_add_inline_runs` has no handler for,
    so the whole placeholder is silently dropped from the paragraph."""
    from src.conversion.markdown_to_docx import markdown_to_docx

    doc = _load(markdown_to_docx("**Client:** <Client Name>"))
    all_text = " ".join(p.text for p in doc.paragraphs)
    assert "<Client Name>" in all_text
